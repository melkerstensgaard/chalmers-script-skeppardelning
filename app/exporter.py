import re
from pathlib import Path
from pypdf import PdfReader, PdfWriter
from .structure import structure_parts
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

def safe(value):
    return re.sub(r"[^0-9A-Za-zÅÄÖåäö._-]+", "_", value or "utan_personnummer")

def export_project(db, output):
    output=Path(output); output.mkdir(parents=True,exist_ok=True)
    groups=[]; current=None
    for row in db.pages():
        classification=row["classification"] or "ogranskad"
        if classification=="nytt_bevis":
            current={"type":"bevis","rows":[row]}; groups.append(current)
        elif classification=="tillhor_foregaende" and current:
            current["rows"].append(row)
        else:
            current={"type":classification,"rows":[row]}; groups.append(current)
    readers={}; index=[]; counters={}
    for group in groups:
        first=group["rows"][0]; relative=Path(first["relative_pdf"]); series,volume=structure_parts(relative)
        key=(str(relative.parent),relative.stem); counters[key]=counters.get(key,0)+1
        filename=f'{relative.stem}_{counters[key]:04d}_{safe(first["personnummer"])}.pdf'
        destination=output/relative.parent; destination.mkdir(parents=True,exist_ok=True); writer=PdfWriter()
        for row in group["rows"]:
            if row["source_pdf"] not in readers: readers[row["source_pdf"]]=PdfReader(row["source_pdf"])
            writer.add_page(readers[row["source_pdf"]].pages[row["page_number"]-1])
        with (destination/filename).open("wb") as file: writer.write(file)
        sources={row["personnummer_source"] for row in group["rows"] if row["personnummer_source"]}
        source="människa" if "människa" in sources else ("skript" if "skript" in sources else "")
        index.append([filename,volume,series,first["personnummer"],group["type"],source])
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Index"

    headers = [
        "Filnamn",
        "Volym",
        "Serie",
        "Personnummer",
        "Dokumenttyp",
        "Personnummer bestämt av",
    ]

    worksheet.append(headers)

    for row in index:
        worksheet.append(row)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    # Behandla personnummer som text så att Excel inte ändrar formatet.
    for cell in worksheet["D"][1:]:
        cell.number_format = "@"

    # Aktivera filter och lås rubrikraden.
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.freeze_panes = "A2"

    # Anpassa kolumnbredderna efter innehållet.
    for column_cells in worksheet.columns:
        max_length = 0

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))

        column_letter = get_column_letter(
            column_cells[0].column
        )

        worksheet.column_dimensions[column_letter].width = min(
            max_length + 2,
            60
        )

    index_path = output / "index.xlsx"
    workbook.save(index_path)
    return len(groups)
