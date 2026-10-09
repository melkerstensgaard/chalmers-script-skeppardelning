import re
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class RegisterRecord:
    personnummer: str
    fornamn: str
    efternamn: str


@dataclass
class NameMatchResult:
    fornamn_similarity: float
    efternamn_similarity: float
    fornamn_level: str
    efternamn_level: str
    points: int
    reasons: list[str]


def normalize_personnummer(value) -> str:
    """
    Normaliserar personnummer till YYMMDD-NNNN.

    Exempel:
    6501011234   -> 650101-1234
    196501011234 -> 650101-1234
    650101-1234  -> 650101-1234
    """

    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    digits = re.sub(
        r"\D",
        "",
        str(value),
    )

    if len(digits) == 12:
        digits = digits[2:]

    if len(digits) != 10:
        return ""

    return f"{digits[:6]}-{digits[6:]}"


def normalize_name(value) -> str:
    """
    Normaliserar namn för jämförelse.

    Exempel:
    Sjöberg   -> sjoberg
    Karl-Erik -> karl erik
    Åström    -> astrom
    """

    if value is None:
        return ""

    text = str(value).casefold().strip()

    text = unicodedata.normalize(
        "NFKD",
        text,
    )

    text = "".join(
        character
        for character in text
        if not unicodedata.combining(character)
    )

    text = re.sub(
        r"[^a-z0-9]+",
        " ",
        text,
    )

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text.strip()


def normalize_header(value) -> str:
    if value is None:
        return ""

    return normalize_name(value).replace(" ", "")


def find_register_columns(worksheet):
    """
    Söker efter rubrikerna Personnummer, Förnamn och Efternamn
    i de första 30 raderna.
    """

    required_headers = {
        "personnummer",
        "fornamn",
        "efternamn",
    }

    maximum_header_row = min(
        worksheet.max_row,
        30,
    )

    for row_number in range(
        1,
        maximum_header_row + 1,
    ):
        current_headers = {}

        for column_number in range(
            1,
            worksheet.max_column + 1,
        ):
            cell_value = worksheet.cell(
                row=row_number,
                column=column_number,
            ).value

            header = normalize_header(
                cell_value,
            )

            if header:
                current_headers[
                    header
                ] = column_number

        if required_headers.issubset(
            current_headers.keys()
        ):
            return {
                "header_row": row_number,
                "personnummer": current_headers[
                    "personnummer"
                ],
                "fornamn": current_headers[
                    "fornamn"
                ],
                "efternamn": current_headers[
                    "efternamn"
                ],
            }

    return None

def load_register(register_path) -> dict[str, list[RegisterRecord]]:
    """
    Läser registret och skapar ett index baserat på personnummer.

    Ett personnummer kan ha flera registerrader. Därför är varje
    värde i resultatet en lista.
    """
    print(f"Läser register: {register_path}")
    path = Path(register_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Registerfilen finns inte: {path}"
        )

    if path.suffix.lower() != ".xlsx":
        raise ValueError(
            "Registerfilen måste vara en XLSX-fil."
        )

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    register_by_personnummer = {}
    found_register_sheet = False

    try:
        for worksheet in workbook.worksheets:
            columns = find_register_columns(
                worksheet,
            )

            if columns is None:
                continue

            found_register_sheet = True

            for row in worksheet.iter_rows(
                    min_row=columns["header_row"] + 1,
                    values_only=True,
            ):

                personnummer_value = row[
                    columns["personnummer"] - 1
                    ]

                personnummer = normalize_personnummer(
                    personnummer_value
                )

                if not personnummer:
                    continue

                fornamn_value = row[
                    columns["fornamn"] - 1
                    ]

                efternamn_value = row[
                    columns["efternamn"] - 1
                    ]

                record = RegisterRecord(
                    personnummer=personnummer,
                    fornamn=(
                        str(fornamn_value).strip()
                        if fornamn_value is not None
                        else ""
                    ),
                    efternamn=(
                        str(efternamn_value).strip()
                        if efternamn_value is not None
                        else ""
                    ),
                )

                register_by_personnummer.setdefault(
                    personnummer,
                    [],
                ).append(record)


    finally:
        workbook.close()

    if not found_register_sheet:
        raise ValueError(
            "Programmet hittade inget kalkylblad med kolumnerna "
            "Personnummer, Förnamn och Efternamn."
        )

    if not register_by_personnummer:
        raise ValueError(
            "Registrets kolumner hittades, men inga giltiga "
            "personnummer kunde läsas."
        )
    print(
        f"Register klart. "
        f"{len(register_by_personnummer)} "
        f"unika personnummer"
    )
    return register_by_personnummer


def sequence_similarity(
    expected_text: str,
    detected_text: str,
) -> float:
    if not expected_text or not detected_text:
        return 0.0

    return SequenceMatcher(
        None,
        expected_text,
        detected_text,
        autojunk=False,
    ).ratio()


def best_name_similarity(
    expected_name: str,
    ocr_text: str,
) -> float:
    """
    Beräknar bästa matchningen mellan ett namn och OCR-texten.

    Först kontrolleras en exakt förekomst. Därefter jämförs
    namnet med ordsekvenser av ungefär samma längd.
    """

    expected = normalize_name(
        expected_name,
    )

    normalized_ocr = normalize_name(
        ocr_text,
    )

    if not expected or not normalized_ocr:
        return 0.0

    if expected in normalized_ocr:
        return 1.0

    expected_words = expected.split()
    ocr_words = normalized_ocr.split()

    if not expected_words or not ocr_words:
        return 0.0

    expected_word_count = len(
        expected_words,
    )

    minimum_window = max(
        1,
        expected_word_count - 1,
    )

    maximum_window = min(
        len(ocr_words),
        expected_word_count + 1,
    )

    best_similarity = 0.0

    for window_size in range(
        minimum_window,
        maximum_window + 1,
    ):
        for start_index in range(
            0,
            len(ocr_words) - window_size + 1,
        ):
            candidate_words = ocr_words[
                start_index:start_index + window_size
            ]

            candidate_text = " ".join(
                candidate_words,
            )

            similarity = sequence_similarity(
                expected,
                candidate_text,
            )

            if similarity > best_similarity:
                best_similarity = similarity

    return best_similarity


def classify_similarity(similarity: float) -> str:
    if similarity >= 0.95:
        return "exact"

    if similarity >= 0.75:
        return "approximate"

    return "missing"


def evaluate_register_record(
    record: RegisterRecord,
    ocr_text: str,
) -> NameMatchResult:
    fornamn_similarity = best_name_similarity(
        record.fornamn,
        ocr_text,
    )

    efternamn_similarity = best_name_similarity(
        record.efternamn,
        ocr_text,
    )

    fornamn_level = classify_similarity(
        fornamn_similarity,
    )

    efternamn_level = classify_similarity(
        efternamn_similarity,
    )

    points = 0
    reasons = []

    if fornamn_level == "exact":
        points += 15
        reasons.append(
            f'Förnamnet "{record.fornamn}" matchar registret: '
            "+15 poäng"
        )

    elif fornamn_level == "approximate":
        points += 8
        reasons.append(
            f'Förnamnet "{record.fornamn}" matchar ungefärligt: '
            "+8 poäng"
        )

    else:
        points -= 10
        reasons.append(
            f'Förnamnet "{record.fornamn}" hittades inte: '
            "-10 poäng"
        )

    if efternamn_level == "exact":
        points += 20
        reasons.append(
            f'Efternamnet "{record.efternamn}" matchar registret: '
            "+20 poäng"
        )

    elif efternamn_level == "approximate":
        points += 10
        reasons.append(
            f'Efternamnet "{record.efternamn}" matchar ungefärligt: '
            "+10 poäng"
        )

    else:
        points -= 15
        reasons.append(
            f'Efternamnet "{record.efternamn}" hittades inte: '
            "-15 poäng"
        )

    return NameMatchResult(
        fornamn_similarity=fornamn_similarity,
        efternamn_similarity=efternamn_similarity,
        fornamn_level=fornamn_level,
        efternamn_level=efternamn_level,
        points=points,
        reasons=reasons,
    )


def find_best_register_record(
    records: list[RegisterRecord],
    ocr_text: str,
):
    """
    Om flera registerrader har samma personnummer väljs den
    registerrad vars namn stämmer bäst med OCR-texten.
    """

    best_record = None
    best_result = None
    best_combined_similarity = -1.0

    for record in records:
        result = evaluate_register_record(
            record,
            ocr_text,
        )

        combined_similarity = (
            result.fornamn_similarity
            + result.efternamn_similarity
        )

        if combined_similarity > best_combined_similarity:
            best_combined_similarity = combined_similarity
            best_record = record
            best_result = result

    return best_record, best_result